"""Build "Review 3 Report.docx" from the Review 2 report + Review 3 evidence.

    python scripts/review3_build_report.py <review2.docx> <out.docx>

Keeps every page of the Review 2 report (same template styles), then:
  - updates the abstract,
  - replaces the ER diagram with docs/review3/ER_DIAGRAM_REVIEW3.png and
    rewrites the text around it,
  - adds the section "Review Requirements (Review 3)" with the checklist,
    demo flow, clean build, concurrency, indexing, testing, working-module
    screenshots and design changes since Review 2,
  - fills the Review 3 evaluation table.
New paragraphs/tables are deep copies of existing ones, so fonts and spacing
match the rest of the report. Needs python-docx.
"""
import copy
import sys
from pathlib import Path

import docx
from docx.shared import Emu, Inches
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "review3" / "screenshots"
ER_PNG = ROOT / "docs" / "review3" / "ER_DIAGRAM_REVIEW3.png"

d = docx.Document(str(SRC))
body = d.element.body


def P(el):
    return Paragraph(el, d)


def find_par(text, style=None, start=0):
    for i, el in enumerate(list(body.iterchildren())):
        if i < start or not el.tag.endswith("}p"):
            continue
        p = P(el)
        if p.text.strip() == text and (style is None or p.style.name == style):
            return el
    raise SystemExit(f"paragraph not found: {text!r}")


def find_par_prefix(prefix):
    for el in body.iterchildren():
        if el.tag.endswith("}p") and P(el).text.startswith(prefix):
            return el
    raise SystemExit(f"paragraph not found: {prefix!r}")


# ---------------------------------------------------------------- prototypes
PROTO = {
    "h1": copy.deepcopy(find_par("Work break down", "Heading 1")),
    "h2": copy.deepcopy(find_par("Design Changes Since Review 1", "Heading 2")),
    "body": copy.deepcopy(find_par_prefix("The ER diagram below shows")),
    "bullet": copy.deepcopy(find_par_prefix("A user OWNS many watchlists")),
    "title": copy.deepcopy(find_par_prefix("S01 – Database server")),
    "caption": copy.deepcopy(find_par("Screenshot S01")),
    "explain": copy.deepcopy(find_par_prefix("Explanation: We use PostgreSQL 18")),
}
# empty paragraph with a little space after it (used after each new table)
PROTO["spacer"] = copy.deepcopy(PROTO["caption"])
for r in PROTO["spacer"].findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"):
    PROTO["spacer"].remove(r)
img_el = find_par("Screenshot S01").getprevious()
PROTO["image"] = copy.deepcopy(img_el)
tables = list(body.iterchildren("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl"))


def table_with(first_row_text):
    for t in tables:
        tb = Table(t, d)
        if tb.rows[0].cells[0].text.strip() == first_row_text[0] and \
                len(tb.columns) == first_row_text[1]:
            return copy.deepcopy(t)
    raise SystemExit(f"table not found {first_row_text}")


TBL = {
    2: table_with(("Change", 2)),
    3: table_with(("#", 3)),
    4: table_with(("Table", 4)),
}


def set_text(p_el, text):
    """Replace a paragraph's text but keep the first run's formatting."""
    p = P(p_el)
    runs = p.runs
    if not runs:
        p.add_run(text)
        return p_el
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    runs[0].text = text
    return p_el


def new_par(kind, text=""):
    el = copy.deepcopy(PROTO[kind])
    if kind == "explain":
        p = P(el)
        label, _, rest = text.partition(":")
        p.runs[0].text = label + ":"
        for r in p.runs[2:]:
            r._element.getparent().remove(r._element)
        p.runs[1].text = " " + rest.strip()
        return el
    return set_text(el, text)


def new_image(path, max_w=6.4, max_h=8.2):
    el = copy.deepcopy(PROTO["image"])
    for r in el.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"):
        el.remove(r)
    w, h = Image.open(path).size
    width = min(max_w, w / 150)
    height = width * h / w
    if height > max_h:
        height = max_h
        width = height * w / h
    P(el).add_run().add_picture(str(path), width=Inches(width))
    return el


def cell_set(cell, text):
    ps = cell.paragraphs
    for extra in ps[1:]:
        extra._element.getparent().remove(extra._element)
    p = ps[0]
    runs = p.runs
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    if not runs:
        p.add_run("")
        runs = p.runs
    lines = text.split("\n")
    runs[0].text = lines[0]
    for line in lines[1:]:
        runs[0].add_break()
        runs[0]._element.append(copy.deepcopy(runs[0]._element.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")))
        runs[0]._element.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")[-1].text = line


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def set_widths(el, inches):
    """Set the column widths (grid + every cell) of a table, in inches."""
    dxa = [str(int(x * 1440)) for x in inches]
    grid = el.find(W_NS + "tblGrid")
    for col, w in zip(grid.findall(W_NS + "gridCol"), dxa):
        col.set(W_NS + "w", w)
    tblw = el.find(W_NS + "tblPr").find(W_NS + "tblW")
    if tblw is not None:
        tblw.set(W_NS + "w", str(sum(int(x) for x in dxa)))
        tblw.set(W_NS + "type", "dxa")
    for tr in el.findall(W_NS + "tr"):
        for tc, w in zip(tr.findall(W_NS + "tc"), dxa):
            tcw = tc.find(W_NS + "tcPr").find(W_NS + "tcW")
            tcw.set(W_NS + "w", w)
            tcw.set(W_NS + "type", "dxa")


def new_table(header, rows, widths=None):
    el = copy.deepcopy(TBL[len(header)])
    t = Table(el, d)
    trs = [r._tr for r in t.rows]
    proto = trs[1]
    for tr in trs[1:]:
        el.remove(tr)
    for c, txt in zip(t.rows[0].cells, header):
        cell_set(c, txt)
    for row in rows:
        tr = copy.deepcopy(proto)
        el.append(tr)
        cells = Table(el, d).rows[-1].cells
        for c, txt in zip(cells, row):
            cell_set(c, txt)
    if widths:
        set_widths(el, widths)
    return el


class Writer:
    """Insert elements one after another before a fixed anchor element."""

    def __init__(self, anchor):
        self.anchor = anchor

    def add(self, el):
        self.anchor.addprevious(el)
        return el

    def h1(self, t): return self.add(new_par("h1", t))
    def h2(self, t): return self.add(new_par("h2", t))
    def text(self, t): return self.add(new_par("body", t))
    def bullet(self, t): return self.add(new_par("bullet", t))
    def table(self, header, rows, widths=None):
        t = self.add(new_table(header, rows, widths))
        self.add(copy.deepcopy(PROTO["spacer"]))
        return t

    def shot(self, title, file, caption, explanation, label="Explanation"):
        self.add(new_par("title", title))
        self.add(new_image(SHOTS / file))
        self.add(new_par("caption", caption))
        if isinstance(explanation, str):
            explanation = [(label, explanation)]
        for lab, txt in explanation:
            self.add(new_par("explain", f"{lab}: {txt}"))


# =========================================================================
# 1. Abstract
set_text(find_par_prefix("In this second review we show"),
         "In this third and final review we present the complete, working system. The database is "
         "fully populated and can be rebuilt from scratch with one command. On top of the Review 2 "
         "work (36 named constraints, 38 SQL queries, 10 complex queries, 3 views, stored procedures, "
         "functions, 2 triggers, cursors, role-based security and an index benchmark), this report "
         "adds a corrected ER diagram, a concurrency test with two database sessions, 221 automated "
         "test checks that all pass, the six-step live demo, and screenshots of every working module "
         "of the web application with its input and output.")

# =========================================================================
# 2. ER diagram: new image + text
er_text = find_par_prefix("The ER diagram below shows")
set_text(er_text,
         "The ER diagram below shows the 6 entity types of our system and the 7 relationships "
         "between them, drawn in Chen notation (rectangle = entity, diamond = relationship, "
         "oval = attribute, underlined = key, double line = every entity on that side must take part). "
         "The diagram was corrected after Review 2; the changes are listed in the table after the "
         "relationship list.")
img_par = er_text.getnext()
blip = img_par.xpath(".//a:blip")[0]
rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
d.part.related_parts[rid]._blob = ER_PNG.read_bytes()
w_px, h_px = Image.open(ER_PNG).size
cx = 6035040
cy = int(cx * h_px / w_px)
for ext in img_par.xpath(".//wp:extent | .//a:ext"):
    if ext.get("cx"):
        ext.set("cx", str(cx))
        ext.set("cy", str(cy))
set_text(img_par.getnext(),
         "Figure 5.1: ER diagram (Chen notation) of the Real-Time Stock Market Watchlist & Alert System")

rel_texts = [
    "A user OWNS many watchlists (1 : N). Every watchlist must have exactly one owner (double line).",
    "A watchlist CONTAINS many instruments, and an instrument can be in many watchlists (M : N). "
    "The relationship has its own attribute added_at. In the database it becomes the table watchlist_items.",
    "An instrument HAS many price ticks (1 : N). Every tick belongs to exactly one instrument.",
    "A user CREATES many alert rules (1 : N), and each alert rule MONITORS exactly one instrument (N : 1).",
    "An alert rule GENERATES many alert events (1 : N), and each alert event is TRIGGERED_BY exactly "
    "one price tick (N : 1).",
]
first_bullet = find_par_prefix("A user OWNS many watchlists")
bullets = [first_bullet]
nxt = first_bullet.getnext()
while nxt is not None and nxt.tag.endswith("}p") and P(nxt).style.name == "Normal" and \
        nxt.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr") is not None:
    bullets.append(nxt)
    nxt = nxt.getnext()
for el, t in zip(bullets, rel_texts):
    set_text(el, t)
for el in bullets[len(rel_texts):]:
    body.remove(el)
w = Writer(bullets[len(rel_texts) - 1].getnext())
w.add(new_par("title", "Corrections made to the ER diagram after Review 2"))
w.table(["Problem in the Review 2 diagram", "What we changed"], [
    ("The lines had arrow heads.",
     "All lines are now plain lines. In Chen notation the direction is not shown with arrows; "
     "the numbers 1, N and M next to each line show the cardinality."),
    ("Watchlist_item was drawn as a weak entity (double rectangle).",
     "It is now the M : N relationship CONTAINS between WATCHLIST and INSTRUMENT, with the attribute "
     "added_at. It is not a weak entity because it has no partial key of its own: a row is identified "
     "only by the watchlist and the instrument it connects."),
    ("Attributes that only join tables were shown, e.g. watchlist_id and instrument_id inside "
     "Watchlist_item.",
     "All foreign-key (join) attributes are removed. In an ER diagram a link between two entities is "
     "shown by a relationship diamond. The foreign-key columns appear only later, when the diagram is "
     "turned into tables (section 6)."),
    ("Some relationship names did not match the tables (ADDED_TO, HAS_RULES).",
     "Relationship names now read as sentences: USER OWNS WATCHLIST, ALERT_RULE MONITORS INSTRUMENT, "
     "ALERT_EVENT TRIGGERED_BY PRICE_TICK."),
])
w.text("ALERT_EVENT is drawn as an entity (not as a relationship) because every alert has its own "
       "number (event_id) that the application uses on its own: the trigger sends it with pg_notify "
       "and the Alert History page lists alerts by it.")

# =========================================================================
# 3. Review 3 section, inserted before "Work break down"
w = Writer(find_par("Work break down", "Heading 1"))
w.h1("Review Requirements (Review 3)")
w.text("This section shows the Review 3 requirements: the complete database, scripts that can be run "
       "again from scratch, transactions and concurrency, indexing and query optimization, testing, the "
       "live demonstration and screenshots of every working module of the project. Everything in the "
       "earlier sections (SQL queries, views, procedures, functions, triggers, cursors and security) is "
       "still part of the final system.")

w.h2("Review 3 Checklist")
w.table(["#", "Requirement", "Where it is shown"], [
    ("1", "Complete database with all tables populated",
     "All 7 tables have data: 3 users, 7 instruments, 5 watchlists, 11 watchlist items, 6 alert rules, "
     "30 price ticks and 8 alert events after the replay (R3_01, R3_02, M04)"),
    ("2", "Full scripts, re-runnable from scratch",
     "scripts/reset_demo_db.sh runs sql/00, 01, 03, 06 and 09 on an empty database (R3_01)"),
    ("3", "Advanced SQL", "38 queries A1–G3 and 10 complex business queries CQ01–CQ10"),
    ("4", "Views", "V1 (4-table join), V2 (aggregate), V3 (row security), V4 (updatability)"),
    ("5", "Procedures", "sp_create_watchlist, sp_set_rule_active, sp_print_watchlist_report (P1–P3, CR3)"),
    ("6", "Functions", "fn_latest_price, fn_price_change_pct, fn_rule_status_report, ingest_tick() (U1, U2, CR1, CR2, S10)"),
    ("7", "Triggers", "trg_price_ticks_evaluate_alerts and trg_alert_events_check_instrument (T1–T5, UI06)"),
    ("8", "Transactions and concurrency", "X01 (SAVEPOINT + COMMIT), X02 (ROLLBACK), P2 (failure → rollback), R3_03 (two sessions)"),
    ("9", "Indexing", "ix_price_ticks_instrument_time (S11, IX1)"),
    ("10", "Query optimization", "Latest 50 ticks on 500,000 rows: 19.966 ms → 0.100 ms (IX2, IX3)"),
    ("11", "Security implementation", "Roles, GRANT, REVOKE and least privilege (SEC1–SEC5, V3)"),
    ("12", "Testing (≥ 15 cases)", "9 test suites, 221 checks, all passed (R3_04 and the table below)"),
    ("13", "Working modules with input and output", "UI01–UI11 (web application)"),
    ("14", "Live demonstration", "Six-step demo flow below"),
])

w.h2("Live Demonstration Flow")
w.text("The live demo follows the six steps asked for in the template. Each step uses a script or "
       "screen that is already shown in this report.")
w.table(["Step", "What we show", "Evidence"], [
    ("1. DDL + DML on a clean database",
     "createdb, then bash scripts/reset_demo_db.sh builds the 7 tables, loads the seed data, creates "
     "the functions, triggers, index, views and roles, and checks the row counts (3|7|5|11|6|0|0). "
     "Then the replay feed inserts 30 ticks and fires 8 alerts.", "R3_01, R3_02"),
    ("2. Role-based access",
     "SET ROLE sw_analyst: reading prices works, but DELETE on price_ticks and SELECT on users give "
     "'permission denied'. arjun sees only his own alerts through the view v_my_alerts.", "SEC5, V3"),
    ("3. Transaction with a deliberate failure and rollback",
     "sp_create_watchlist is called with one unknown symbol. The procedure raises an error, so the "
     "watchlist and its items are not created at all (all or nothing).", "P2, X02"),
    ("4. Trigger firing",
     "HDFCBANK rule ABOVE 1700: a tick at 1690 does nothing, a tick at 1710 crosses 1700 and the "
     "trigger inserts one row into alert_events.", "T2, UI05, UI06"),
    ("5. One optimization before / after",
     "Latest 50 RELIANCE ticks on 500,000 rows: without the index PostgreSQL reads about 25,000 rows "
     "and sorts them (19.966 ms); with the index it reads 50 rows in order (0.100 ms).", "IX2, IX3"),
    ("6. One concurrency scenario",
     "Two sessions send RELIANCE prices at the same time. The second session waits for the first "
     "one's row lock, so only one alert is created.", "R3_03"),
], widths=[1.65, 4.05, 0.95])

w.h2("Building the Database From Scratch")
w.shot("R3_01 – Demo step 1: DDL + DML on a new, empty database",
       "R3_01_clean_build.png", "Screenshot R3_01",
       "A new empty database (stock_watchlist_review3) is created and the reset script runs all the "
       "SQL scripts in order: tables and constraints, seed data, ingest_tick() and the triggers, the "
       "index, and the views, procedures and roles. The last line checks the row counts: 3 users, "
       "7 instruments, 5 watchlists, 11 items, 6 rules, 0 ticks and 0 alerts. The same script can be "
       "run again at any time and always gives the same clean state.")
w.shot("R3_02 – Loading prices with the replay feed",
       "R3_02_replay.png", "Screenshot R3_02",
       "The replay feed reads 30 prices from data/replay_prices.csv and sends each one to "
       "ingest_tick(). All 30 were inserted, with 0 duplicates and 0 errors. The trigger fired 8 "
       "alerts while the data was loading (the ALERT lines). After this step every table has data.")

w.h2("Transactions and Concurrency")
w.text("A transaction is a group of SQL statements that either all succeed (COMMIT) or all have no "
       "effect (ROLLBACK). COMMIT, ROLLBACK and SAVEPOINT are shown in X01, X02 and P2. Our database "
       "runs at PostgreSQL's default isolation level, READ COMMITTED: each statement sees only data "
       "that other sessions have already committed.")
w.text("Concurrency means two or more sessions using the database at the same time. The risky case in "
       "our project is two feeds sending prices for the same instrument at the same moment. To compare "
       "the previous price with the new one safely, ingest_tick() first locks the instrument's row "
       "(SELECT ... FOR NO KEY UPDATE). A second session that wants the same instrument has to wait "
       "until the first one commits. Different instruments do not block each other.")
w.shot("R3_03 – Demo step 6: two sessions send prices for the same instrument",
       "R3_03_concurrency.png", "Screenshot R3_03", [
           ("Scenario 1",
            "The last committed RELIANCE price is 2990 and rule #1 is ABOVE 3000. Session A sends 3005 "
            "(this crosses 3000, so one alert fires) and keeps its transaction open for 3 seconds. One "
            "second later session B sends 3010. A third session (observer) sees in pg_stat_activity "
            "that B is waiting for a Lock held by A. When A commits, B continues after waiting 2.02 s. "
            "B now sees 3005 as the previous price, so 3005 → 3010 is not a new crossing and no second "
            "alert is made. Result: exactly 1 alert for 1 real crossing."),
           ("Without the lock",
            "B would not wait. It would still see 2990 as the previous price (A had not committed yet), "
            "so 2990 → 3010 would also look like a crossing and the user would get the same alert twice."),
           ("Scenario 2",
            "Both sessions send the same price event (source_event_id 'cc:dup') at the same time. One "
            "is INSERTED and the other gets DUPLICATE, because of UNIQUE (source, instrument_id, "
            "source_event_id). Only 1 row is stored."),
       ])
w.text("The same behaviour is checked automatically by tests/concurrency_test.sh (17 checks with "
       "separate psql sessions), which is part of the test run in R3_04.")

w.h2("Indexing and Query Optimization")
w.text("The most common search in our project is 'the latest prices of one instrument'. It is used by "
       "the dashboard, the watchlist page, the instrument history page and by the alert trigger, which "
       "needs the previous price for every new tick. For this search we created one B-tree index:")
w.bullet("ix_price_ticks_instrument_time ON price_ticks (instrument_id, observed_at DESC, tick_id DESC)")
w.text("instrument_id comes first because every one of these queries filters on one instrument. "
       "observed_at DESC and tick_id DESC come next, so the newest ticks are already at the front of "
       "the index in the right order and PostgreSQL does not need a Sort step. The instrument search by "
       "exchange and symbol (used by ingest_tick()) is fast because of the index that PostgreSQL "
       "creates for UNIQUE (exchange, symbol).")
w.text("The before/after comparison is shown in IX2 and IX3 (section 'Indexes and EXPLAIN Comparison'). "
       "On a test copy with 500,000 ticks, the query for the latest 50 RELIANCE ticks went from a Bitmap "
       "Heap Scan + Sort reading about 25,000 rows in 19.966 ms to an Index Scan reading 50 rows in "
       "0.100 ms. The cost is extra disk space (19 MB) and a little extra work on every INSERT. On the "
       "small demo data (30 ticks) the index makes no visible difference, which is why the test was "
       "done on 500,000 rows.")

w.h2("Testing")
w.text("We test the project with 9 automated test suites. The SQL suites run inside a transaction that "
       "is rolled back, and the Python suites use temporary databases, so the demo database is never "
       "changed. In total there are 221 checks plus the schema verifier, and all of them pass "
       "(screenshot R3_04). The table lists 20 of the test cases.")
w.table(["Test case", "Expected result", "Result"], [
    ("A1 duplicate username", "Rejected by uq_users_username", "PASS"),
    ("A3 duplicate instrument (exchange, symbol)", "Rejected by UNIQUE", "PASS"),
    ("A8 duplicate tick (source, instrument, event id)", "Rejected by UNIQUE", "PASS"),
    ("B1 negative price", "Rejected by CHECK", "PASS"),
    ("B4 NULL (unknown) volume", "Accepted (volume may be unknown)", "PASS"),
    ("C1 watchlist for a user that does not exist", "Rejected by FOREIGN KEY", "PASS"),
    ("D3 delete an instrument that has price history", "Blocked by ON DELETE RESTRICT", "PASS"),
    ("E3 delete a user", "Their watchlists, rules and alerts are deleted (CASCADE)", "PASS"),
    ("E7 delete a user", "All price ticks are kept", "PASS"),
    ("01c first tick of an instrument", "No alert (no previous price)", "PASS"),
    ("02b price 2999 → 3001 with rule ABOVE 3000", "Exactly one alert", "PASS"),
    ("03 price stays above 3001 → 3010 → 3020", "Still only one alert", "PASS"),
    ("09a second crossing inside the cooldown", "No alert", "PASS"),
    ("13c late (out-of-order) tick above the threshold", "Stored, but no alert", "PASS"),
    ("17a ingest_tick with an unknown symbol", "Error SW001", "PASS"),
    ("18d rollback to a savepoint", "Tick and alert disappear together", "PASS"),
    ("Concurrency: two sessions, same instrument", "Second session waits; one alert", "PASS"),
    ("API test_10 duplicate watchlist name", "HTTP 409 with a readable message", "PASS"),
    ("API test_29 SQL-injection text as input", "Stored as plain text, no harm", "PASS"),
    ("Index test_07 latest-N query", "Uses the index, no Sort step", "PASS"),
], widths=[3.1, 2.75, 0.8])
w.shot("R3_04 – All test suites", "R3_04_all_tests.png", "Screenshot R3_04",
       "scripts/run_all_tests.sh runs all 9 suites one after another and stops at the first failure. "
       "Every suite passed: schema tests 51/51, schema verifier, alert tests 50/50, concurrency 17, "
       "replay 18, index 15, API 32, frontend 14 and live feed 24. (This run was on Windows, where the "
       "two SQL test files send their hidden output to NUL instead of /dev/null.)")

w.h2("Working Modules of the Project (Web Application)")
w.text("The web application has six modules: Dashboard, Watchlists, Alert Rules, Alert History, "
       "Instruments and Demo (manual tick). The browser only sends requests to FastAPI; all the data "
       "rules (keys, checks, the alert trigger) run inside PostgreSQL. The screenshots below were taken "
       "in this order as user arjun, right after the replay in R3_02.")
w.shot("UI01 – Dashboard", "UI01_dashboard.png", "Screenshot UI01", [
    ("Input", "Choose the demo user (arjun) at the top. 'Refresh' and 'Live refresh' re-read the data."),
    ("Output", "Counts of watchlists, active rules, recent alerts and instruments, the 5 latest alerts, "
               "and every watchlist with the latest price of each instrument (HDFCBANK has no price yet)."),
])
w.shot("UI02 – Watchlists: create a watchlist and add instruments", "UI02_watchlists.png",
       "Screenshot UI02", [
           ("Input", "Name 'Banking' → Create watchlist; then 'Add instrument' HDFCBANK and TCS."),
           ("Output", "The new watchlist with its 2 instruments and their latest prices. Each list can "
                      "be deleted, and each instrument can be removed."),
       ])
w.shot("UI03 – Watchlists: a duplicate name is rejected", "UI03_watchlist_duplicate_error.png",
       "Screenshot UI03", [
           ("Input", "Create another watchlist called 'Crypto' for arjun, who already has one."),
           ("Output", "The message 'arjun already has a watchlist named \"Crypto\"'. It comes from the "
                      "UNIQUE (user_id, name) constraint in PostgreSQL. The API returns it as HTTP 409 and "
                      "the page shows it as a readable message."),
       ])
w.shot("UI04 – Alert Rules: create a rule", "UI04_alert_rules.png", "Screenshot UI04", [
    ("Input", "Instrument NSE · HDFCBANK, direction ABOVE, threshold 1700, cooldown 0 → Create rule."),
    ("Output", "Rule #7 appears in the list as ACTIVE. Each rule can be disabled/enabled, its cooldown "
               "changed, or deleted. Instrument, direction and threshold cannot be edited (disable and "
               "create a new rule instead)."),
])
w.shot("UI05 – Demo: manual tick that does not cross", "UI05_manual_tick_no_alert.png",
       "Screenshot UI05", [
           ("Input", "Pick HDFCBANK, price 1690, volume 100 → Send tick."),
           ("Output", "Result INSERTED with its tick_id. This is the first HDFCBANK price, so there is "
                      "no previous price and no alert."),
       ])
w.shot("UI06 – Demo: manual tick that fires an alert (trigger)", "UI06_manual_tick_alert_fired.png",
       "Screenshot UI06", [
           ("Input", "Same instrument, price 1710, volume 150 → Send tick."),
           ("Output", "Tick 32 INSERTED. The price moved from 1690 to 1710 and crossed 1700, so the "
                      "trigger recorded an alert for rule #7. The page shows the new alert row."),
       ])
w.shot("UI07 – Demo: the same tick sent again", "UI07_manual_tick_duplicate.png",
       "Screenshot UI07", [
           ("Input", "Resend last tick (same source_event_id)."),
           ("Output", "Result DUPLICATE, nothing stored, no second alert."),
       ])
w.shot("UI08 – Alert History", "UI08_alert_history.png", "Screenshot UI08", [
    ("Input", "Open Alert History; optionally filter by instrument."),
    ("Output", "arjun's alerts from alert_events joined with the rule and the tick: time fired, "
               "instrument, direction, threshold, the price that triggered it and the rule number, "
               "newest first."),
])
w.shot("UI09 – Instruments: search, latest price and price history",
       "UI09_instrument_search_history.png", "Screenshot UI09", [
           ("Input", "Type 'reli' in Search, then click 'Latest price & history' for RELIANCE."),
           ("Output", "The latest price (2,788.00 INR), a price chart and the list of ticks, newest "
                      "first. This page uses the latest-price and history queries that the index "
                      "ix_price_ticks_instrument_time speeds up."),
       ])
w.shot("UI10 – Dashboard after the demo", "UI10_dashboard_after.png", "Screenshot UI10", [
    ("Input", "Return to Dashboard and press Refresh."),
    ("Output", "3 watchlists (with Banking), 4 active rules and the new HDFCBANK alert at the top."),
])
w.shot("UI11 – API documentation", "UI11_api_docs.png", "Screenshot UI11", [
    ("Input", "Open http://127.0.0.1:8000/docs."),
    ("Output", "FastAPI's automatic page listing all endpoints. Each one can be tried from the browser."),
])

w.h2("Design Changes Since Review 2")
w.table(["Change", "Reason"], [
    ("ER diagram redrawn: plain lines without arrows; Watchlist_item replaced by the M : N relationship "
     "CONTAINS (with added_at); join-only (foreign-key) attributes removed.",
     "Review 2 feedback. The new diagram follows Chen notation and shows only real attributes of each "
     "entity."),
    ("Added Review 3 evidence: clean build, concurrency demo (scripts/review3_concurrency_demo.py), "
     "full test run and screenshots of all web modules.",
     "Review 3 asks for concurrency, testing (≥ 15 cases), a live demo and screenshots of all working "
     "modules."),
    ("No change to the tables, constraints, functions or triggers.",
     "The 7 tables were already in BCNF with 36 named constraints; the ER fixes are drawing corrections, "
     "not database changes."),
])

# =========================================================================
# 4. Review evaluation table (Review 3 template rows)
for t in d.tables:
    if t.rows[0].cells[0].text.strip() == "COMPONENT":
        el = t._tbl
        proto = t.rows[-1]._tr
        # rebuild the rows from the existing TOTAL row (keeps the 5-column layout)
        for tr in [r._tr for r in t.rows][1:]:
            el.remove(tr)
        for row in [("Project work with screenshot", "5"), ("Viva", "5"), ("TOTAL", "10 MARKS")]:
            tr = copy.deepcopy(proto)
            el.append(tr)
            cells = Table(el, d).rows[-1].cells
            cell_set(cells[0], row[0])
            cell_set(cells[1], row[1])
        break

d.save(str(OUT))
print("wrote", OUT)
